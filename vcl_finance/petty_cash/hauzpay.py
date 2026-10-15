"""Hauz Pay statement import — the Frappe layer over ``hauzpay_core``.

Someone downloads the weekly Hauz Pay statement (.xlsx), uploads it (a private File),
and calls:

* ``preview_statement(file_url)`` — what WOULD happen, per Sun-Sat week. Writes nothing.
* ``import_statement(file_url, week_ending=None)`` — writes the new lines as ordinary
  rows on that week's Hauz-Pay ``Petty Cash Sheet`` (created if missing). The sheet save
  mirrors them into ``Petty Cash Entry`` as Unapproved, so they land in the normal
  approval queue.

THE STATEMENT SETS THE OPENING (Tanuj, 15/10/2026), as a bank statement would: each
week's sheet opens at the statement's balance at the start of that week, on a sheet the
import creates or one still open, whatever the carry-forward said (a sheet's opening is
otherwise a snapshot taken when it was created, and nothing updates it). The old value
and the difference are in the preview and the result. An Approved / Closed / Submitted
sheet is never changed, only flagged. See ``hauzpay_core.decide_opening``.

Nothing here posts to ERPNext or QuickBooks; posting stays with ``weekpost``.

NOTHING IS DECIDED HERE (Tanuj, 14/10/2026). The route, category and bank the rules
find are SUGGESTIONS, written into the row's notes as "Suggested: ..." for the reviewer
to accept or change on the normal approval screen, the same review as for a line the
custodian keys by hand. The import never locks a row (the lock IS the approval), never
sets a ``Petty Cash Entry``'s ``posting_account`` (a line-level account is a human
decision: it beats the Posting Map in ``post._classify``), and never posts or stages
anything. ``_assert_nothing_decided`` checks that on every write and rolls the whole
import back if a mirrored entry comes out anything but Unapproved and undecided.

Every imported row carries its Hauz Pay transaction id in ``src_ref`` (the one lumped
charges row per week carries ``HZP-CHARGES`` plus the id of every charge folded into
it). An id found on ANY Hauz-Pay sheet, cancelled rows included, is never imported
again, so re-uploading an overlapping statement is safe. Closed / Submitted /
Approved weeks are never touched.

Rules (roster, known payees, keywords) live in ``hauzpay_rules.json`` beside this file
— a file rather than a doctype for now: they change rarely and go through review with
the code. On top of the roster, a phone that was already paid on the WAGES register of
a Hauz-Pay sheet (recipient written ``254XXXXXXXXX - NAME``) is routed as a wage again.

The whole import is one request, so one transaction: if any week fails, nothing is
written.
"""
import io
import json
import os
import re
from datetime import timedelta

import frappe
from frappe import _
from frappe.utils import flt, getdate

from vcl_finance.petty_cash import hauzpay_core as C
from vcl_finance.petty_cash.api import PETTY_APPROVERS, _first_blank, _next_idx, _summary, _voucher_has_data
from vcl_finance.petty_cash.doctype.petty_cash_sheet.petty_cash_sheet import (
    _prior_close, create_for_week, week_saturday,
)

FLOAT = "Hauz-Pay"
COMPANY = "Vimit Converters Limited"
ERP_ONLY = "ERPNext only, not QuickBooks"
RULES_PATH = os.path.join(os.path.dirname(__file__), "hauzpay_rules.json")
_PHONE_IN_TEXT = re.compile(r"(?:\+?254|\b0)(\d{9})\b")
_WAGE_RECIPIENT = re.compile(r"^\s*\+?254(\d{9})\s*-\s*(.+?)\s*$")


def load_rules():
    with open(RULES_PATH) as f:
        return json.load(f)


def _assert_finance():
    """Same role set as approving a line (``api.PETTY_APPROVERS``), checked in the body
    because a whitelisted method is reachable whatever the doctype permissions say."""
    if not (set(frappe.get_roles()) & PETTY_APPROVERS):
        frappe.throw(_("Only Finance can import a Hauz Pay statement."), frappe.PermissionError)


# ------------------------------------------------------------------ reading
def _read_statement(file_url):
    if not file_url:
        frappe.throw(_("Upload the Hauz Pay statement first."))
    name = frappe.db.get_value("File", {"file_url": file_url}, "name", order_by="creation desc")
    if not name:
        frappe.throw(_("Uploaded file not found: {0}").format(file_url))
    content = frappe.get_doc("File", name).get_content()
    if isinstance(content, str):
        content = content.encode("latin-1")
    from openpyxl import load_workbook
    try:
        # NOT read_only: the Hauz Pay export carries a wrong <dimension> tag and read_only
        # mode trusts it, returning only the first cell.
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        frappe.throw(_("That file is not an Excel (.xlsx) workbook. Download the statement "
                       "from Hauz Pay as Excel and upload it unchanged."))
    ws = wb["Statement"] if "Statement" in wb.sheetnames else wb[wb.sheetnames[0]]
    try:
        return C.parse_rows(ws.iter_rows(values_only=True))
    except ValueError as e:
        frappe.throw(str(e))


def _phone_of(*texts):
    for t in texts:
        m = _PHONE_IN_TEXT.search(str(t or ""))
        if m:
            return m.group(1)
    return ""


def _is_charges_row(r):
    """A lumped-charges voucher keyed before this import existed. Three spellings live:
    voucher_no CHARGES (Jul-Aug backfill), recipient '... transaction charges' (Gen-1),
    recipient 'M-Pesa' with 'Transaction charges (lumped ...)' in the notes (June)."""
    return ((r.voucher_no or "").strip().upper() == "CHARGES"
            or "charges" in (r.recipient or "").lower()
            or "transaction charges" in (r.notes or "").lower())


def _sheet_rows(sheet_name):
    """The sheet's money rows as plain dicts for the core's legacy match."""
    out = []
    for r in frappe.get_all("Petty Cash Voucher", filters={"parent": sheet_name, "parenttype": "Petty Cash Sheet"},
                            fields=["txn_date", "amount", "cash_in", "voucher_no", "recipient",
                                    "notes", "cancelled", "src_ref"]):
        if r.cancelled or not r.txn_date or not flt(r.amount):
            continue
        out.append({"date": getdate(r.txn_date), "amount": flt(r.amount, 2), "cash_in": bool(r.cash_in),
                    "phone": _phone_of(r.voucher_no, r.recipient),
                    "is_charges": _is_charges_row(r),
                    "src_ref": r.src_ref, "label": r.recipient})
    for r in frappe.get_all("Petty Cash Wages Entry", filters={"parent": sheet_name, "parenttype": "Petty Cash Sheet"},
                            fields=["txn_date", "amount", "recipient", "cancelled", "src_ref"]):
        if r.cancelled or not r.txn_date or not flt(r.amount):
            continue
        out.append({"date": getdate(r.txn_date), "amount": flt(r.amount, 2), "cash_in": False,
                    "phone": _phone_of(r.recipient), "is_charges": False,
                    "src_ref": r.src_ref, "label": r.recipient})
    return out


def _live_state(weeks):
    """(sheets by Saturday, known src_refs, learned wage roster, default custodian)."""
    sheets = {}
    all_sheets = frappe.get_all("Petty Cash Sheet", filters={"float": FLOAT, "docstatus": ("<", 2)},
                                fields=["name", "week_ending", "status", "opening_balance", "custodian_name"],
                                order_by="week_ending desc")
    for s in all_sheets:
        sat = week_saturday(getdate(s.week_ending))
        if sat in weeks and sat not in sheets:   # newest wins if a legacy Fri sheet shares the week
            sheets[sat] = {"name": s.name, "status": s.status or "Draft",
                           "opening_balance": flt(s.opening_balance, 2), "rows": _sheet_rows(s.name)}

    known = {}
    for tbl in ("Petty Cash Voucher", "Petty Cash Wages Entry"):
        for parent, ref in frappe.db.sql(f"""
                select c.parent, c.src_ref from `tab{tbl}` c
                join `tabPetty Cash Sheet` s on s.name = c.parent
                where c.parenttype = 'Petty Cash Sheet' and s.float = %s and s.docstatus < 2
                  and ifnull(c.src_ref, '') != ''""", (FLOAT,)):
            for tok in C.ref_tokens(ref):
                known[tok] = parent

    learned = {}
    for (rec,) in frappe.db.sql("""
            select w.recipient from `tabPetty Cash Wages Entry` w
            join `tabPetty Cash Sheet` s on s.name = w.parent
            where w.parenttype = 'Petty Cash Sheet' and s.float = %s and s.docstatus < 2
              and ifnull(w.cancelled, 0) = 0 and w.recipient like '%%-%%'
            order by s.week_ending asc""", (FLOAT,)):
        m = _WAGE_RECIPIENT.match(rec or "")
        if m:
            learned["254" + m.group(1)] = m.group(2)

    custodian = all_sheets[0].custodian_name if all_sheets else None
    return sheets, known, learned, custodian or "Shiro"


def _build_plan(file_url):
    parsed = _read_statement(file_url)
    weeks = {week_saturday(l["date"]) for l in parsed["lines"]}
    sheets, known, learned, custodian = _live_state(weeks)
    p = C.plan(parsed, load_rules(), sheets, learned_roster=learned, known_refs=known)
    _check_banks(p)
    _predict_openings(p)
    return p, custodian


def _check_banks(p):
    """A top-up's suggested bank comes from the rules file; only a LIVE, enabled Bank
    account of the company counts. Anything else is treated as "bank?" rather than
    suggesting a name that does not exist."""
    live = set(frappe.get_all("Account", filters={"account_type": "Bank", "company": COMPANY,
                                                   "is_group": 0, "disabled": 0}, pluck="name"))
    p["bank_accounts"] = sorted(live)
    for w in p["weeks"]:
        for l in w["lines"]:
            acct = l.get("bank_account")
            if l["route"] != C.CASH_IN or not acct or acct in live:
                continue
            l["bank_account"], l["needs_bank"] = None, True
            l["why"] = f"bank? — the rules name {acct}, which is not a live bank account"
            l["suggestion"] = C.suggestion(l)
            p.setdefault("warnings", []).append(
                f"hauzpay_rules.json maps a top-up to '{acct}', which is not an enabled Bank account.")
        w["needs_bank"] = sum(1 for l in w["lines"] if l["state"] == "new" and l.get("needs_bank"))
    p["totals"]["needs_bank"] = sum(w["needs_bank"] for w in p["weeks"])


def _predict_openings(p):
    """Say, per week, what the import does to the sheet's opening (``C.decide_opening``).

    THE STATEMENT IS THE AUTHORITY (Tanuj, 15/10/2026): an open or new sheet's opening is
    set to the statement's balance at the start of its week, whatever the carry-forward
    said. For a week the import will CREATE, the "system" figure is what ``create_for_week``
    would carry forward: the previous week's statement close when that week is in this
    statement too (the import fills it first), otherwise the last sheet on file.
    """
    starts = C.period_start(p.get("period"))
    prev = None
    for w in p["weeks"]:
        if w["action"] == "create":
            system = None
            if prev and prev["action"] != "skip" and \
                    getdate(prev["week_ending"]) == getdate(w["week_ending"]) - timedelta(days=7) \
                    and prev["statement_closing"] is not None:
                system, w["opening_basis"] = prev["statement_closing"], "previous week in this statement"
            else:
                info = _prior_close(FLOAT, w["week_ending"])
                if info:
                    system, w["opening_basis"] = info["balance"], "carried from " + info["prior_name"]
            w["sheet_opening"] = system
            w.update(C.decide_opening(w, system, starts))
        elif w["sheet"]:
            w["opening_basis"] = "sheet " + w["sheet"]
        if w.get("opening_action") == "flag":
            p.setdefault("warnings", []).append(
                f"Week ending {getdate(w['week_ending']).strftime('%d/%m/%Y')}: {w['opening_note']}")
        prev = w


# ------------------------------------------------------------------ API
@frappe.whitelist()
def preview_statement(file_url):
    """Parse the uploaded statement and say, per week, what an import would do. Writes nothing."""
    _assert_finance()
    p, _custodian = _build_plan(file_url)
    p["file_url"] = file_url
    return p


@frappe.whitelist(methods=["POST"])
def import_statement(file_url, week_ending=None):
    """Write the statement's NEW lines onto the Hauz-Pay sheets as ordinary draft rows.

    The plan is rebuilt here from the live data — the preview the browser holds is never
    trusted. ``week_ending`` (any date in the week) limits the import to that one week.
    """
    _assert_finance()
    p, custodian = _build_plan(file_url)
    only = week_saturday(getdate(week_ending)).isoformat() if week_ending else None

    done = []
    for wk in p["weeks"]:
        if only and wk["week_ending"] != only:
            continue
        new = [l for l in wk["lines"] if l["state"] == "new"]
        res = {"week_ending": wk["week_ending"], "sheet": wk["sheet"], "created": False,
               "vouchers": 0, "wages": 0, "charges": 0, "charges_amount": 0.0, "flagged": 0,
               "skipped": len(wk["lines"]) - len(new), "status": wk["status"]}
        if wk["action"] == "skip" or not new:
            res["note"] = (_("Week is {0} — left untouched.").format(wk["status"])
                           if wk["action"] == "skip" else _("Nothing new for this week."))
            res["opening"] = _opening_result(wk)
            if not wk["sheet"]:
                res["opening"].update(action="none", note=_("No sheet for this week and nothing to import, so none was created."))
            elif wk.get("opening_action") == "set" and wk.get("opening_difference"):
                res["opening"] = _set_opening_only(wk)
            elif wk.get("opening_action") == "flag":
                res["warnings"] = [res["opening"]["note"]]
            done.append(res)
            continue
        res.update(_write_week(wk, new, custodian))
        done.append(res)

    totals = {k: sum(d.get(k) or 0 for d in done)
              for k in ("vouchers", "wages", "charges", "flagged", "skipped", "topups", "topups_amount", "needs_bank")}
    warnings = [w for d in done for w in (d.get("warnings") or [])]
    totals["unapproved"] = sum(d.get("unapproved") or 0 for d in done)
    totals["posted"] = 0
    return {"weeks": done, "totals": totals, "balance": p["balance"], "warnings": warnings}


def _assert_nothing_decided(fresh, reused=()):
    """Every entry the import's rows mirror into must still be waiting for a person.

    ``fresh``: child rows this import filled or added. Their entries must be Unapproved,
    with no posting_account, journal or QuickBooks queue row - anything else means
    something decided on the reviewer's behalf, and the whole import is rolled back.
    ``reused``: the week's existing lumped-charges row, topped up with new charges. It
    is unlocked (a locked one is never reused), so it cannot be Approved; a reviewer may
    already have Withdrawn it or picked its account, which the import leaves alone.

    The mirror runs synchronously on the sheet save but swallows its own errors, so a
    row with no entry is reported as a warning, never assumed. Returns
    (fresh rows whose entry is waiting Unapproved, warnings)."""
    rows = [r for r in list(fresh) + list(reused) if r.name]
    if not rows:
        return 0, []
    ents = {e.origin_row: e for e in frappe.get_all(
        "Petty Cash Entry",
        filters={"origin_row": ("in", [r.name for r in rows]), "sync_state": "Mirrored"},
        fields=["name", "origin_row", "status", "posting_account", "journal_entry", "qbo_queue"])}
    bad, warnings = [], []
    reused_names = {r.name for r in reused}
    for r in rows:
        e = ents.get(r.name)
        if not e:
            warnings.append(f"Row {r.idx} ({r.src_ref or r.name}) has no Petty Cash Entry yet; "
                            "it shows in the approval queue once the sheet is next saved.")
        elif r.name in reused_names:
            if e.status == "Approved" or e.journal_entry:
                bad.append(e.name)
        elif e.status != "Unapproved" or e.posting_account or e.journal_entry or e.qbo_queue:
            bad.append(e.name)
    if bad:
        frappe.throw(_("The import would have left {0} approved, posted or with an account "
                       "chosen. Nothing was imported.").format(", ".join(bad)))
    return sum(1 for r in fresh if r.name in ents), warnings


def _write_week(wk, new, custodian):
    we = getdate(wk["week_ending"])
    created = False
    name = wk["sheet"]
    if not name:
        name = create_for_week(week_ending=we, float_name=FLOAT, custodian_name=custodian,
                               opening_balance=wk["statement_opening"] or 0)
        created = True
    doc = _open_sheet(name, we)
    opening = _apply_opening(doc, wk)

    vouchers = sorted(doc.vouchers, key=lambda x: x.row_idx or 0)
    wages = sorted(doc.wages_entries, key=lambda x: x.row_idx or 0)

    # Never a locked row (``_first_blank`` skips them: the lock IS the approval) and
    # never a voided blank, or the imported line would land Approved or Void.
    def voucher_row():
        r = _first_blank(vouchers, lambda v: not _voucher_has_data(v) and not v.cancelled)
        if r is None:
            r = doc.append("vouchers", {"row_idx": _next_idx(doc.vouchers)})
            vouchers.append(r)
        return r

    def wage_row():
        r = _first_blank(wages, lambda w: not w.cancelled and not (w.recipient or "").strip() and not flt(w.amount)
                         and not w.txn_date and not (w.reason or "").strip() and not (w.staff_id or "").strip())
        if r is None:
            r = doc.append("wages_entries", {"row_idx": _next_idx(doc.wages_entries)})
            wages.append(r)
        return r

    out = {"vouchers": 0, "wages": 0, "charges": 0, "charges_amount": 0.0, "flagged": 0,
           "topups": 0, "topups_amount": 0.0, "needs_bank": 0, "warnings": []}
    fresh, reused = [], []
    charges = [l for l in new if l["route"] == C.CHARGE]
    for l in new:
        if l["route"] == C.CHARGE:
            continue
        d = getdate(l["date"])
        if l["kind"] == "wage":
            r = wage_row()
            # A wages row has no notes; the suggestion leads its reason so it survives the 140 cut.
            r.txn_date, r.entry_type, r.recipient = d, "Wage", l["recipient"][:140]
            r.reason = " · ".join([l["suggestion"], l["reason"] or ""]).strip(" ·")[:140]
            r.amount, r.src_ref = l["amount"], l["txid"]
            fresh.append(r)
            out["wages"] += 1
            continue
        r = voucher_row()
        notes = [l["suggestion"]]
        if l["narration"]:
            notes.append(l["narration"])
        if l["expense_category"]:
            notes.append("Hauz Pay category: " + l["expense_category"])
        if l["route"] == C.FLAGGED:
            out["flagged"] += 1
        if l["route"] == C.CASH_IN:
            notes = [f"Top-up · {ERP_ONLY}", f"{l.get('source') or ''} ({l['payee'] or l['name'] or 'no payee'})"] + notes
            out["topups"] += 1
            out["topups_amount"] = round(out["topups_amount"] + flt(l["amount"]), 2)
            if not l.get("bank_account"):
                out["needs_bank"] += 1
        fresh.append(r)
        r.txn_date, r.recipient, r.amount = d, l["recipient"][:140], l["amount"]
        r.voucher_no = (l["txid"] or "")[:20]
        r.cash_in = 1 if l["route"] == C.CASH_IN else 0
        r.category = "" if r.cash_in else (l["category"] or "OT")
        r.notes = " · ".join(notes)
        r.src_ref = l["txid"]
        out["vouchers"] += 1

    if charges:
        amt = round(sum(l["amount"] for l in charges), 2)
        refs = [l["txid"] for l in charges if l["txid"]]
        last = max(getdate(l["date"]) for l in charges)
        row = next((v for v in vouchers if (v.src_ref or "").startswith(C.CHARGES_REF)
                    and not v.locked and not v.cancelled), None)
        if row is None:
            row = voucher_row()
            row.recipient, row.category, row.cash_in, row.amount = "Hauz Pay transaction charges", "OT", 0, 0
            row.voucher_no, row.src_ref = "CHARGES", C.CHARGES_REF
            fresh.append(row)
        else:
            reused.append(row)
        row.amount = round(flt(row.amount) + amt, 2)
        row.txn_date = max(getdate(row.txn_date), last) if row.txn_date else last
        row.src_ref = "\n".join([row.src_ref] + refs)
        n = len([t for t in row.src_ref.split("\n")[1:] if t])
        row.notes = (f"{C.suggestion(charges[0])} · "
                     f"Lumped Hauz Pay transaction charges for the week ({n} charges)")
        out["charges"], out["charges_amount"] = len(charges), amt

    doc.save()
    out["unapproved"], missing = _assert_nothing_decided(fresh, reused)
    out["warnings"].extend(missing)
    out.update({"sheet": doc.name, "created": created, "opening": opening, "summary": _summary(doc)})
    return out


def _open_sheet(name, we):
    doc = frappe.get_doc("Petty Cash Sheet", name)
    if not frappe.has_permission("Petty Cash Sheet", "write", doc=doc):
        frappe.throw(_("You may not write to the petty cash sheet {0}.").format(name), frappe.PermissionError)
    if doc.docstatus != 0 or doc.status in C.LOCKED_STATUSES:
        frappe.throw(_("The Hauz-Pay week ending {0} is {1}; nothing was imported.").format(
            we.strftime("%d/%m/%Y"), doc.status))
    return doc


def _opening_result(wk, old=None, new=None):
    """What the import did to one sheet's opening, for the result. ``old`` is the value the
    sheet actually had (a new sheet: what create_for_week carried forward)."""
    old = wk.get("opening_system") if old is None else old
    return {"action": wk.get("opening_action") or "none", "old": old,
            "new": new if new is not None else old, "statement": wk.get("opening_statement"),
            "difference": wk.get("opening_difference"), "note": wk.get("opening_note") or ""}


def _apply_opening(doc, wk):
    """Set the (open) sheet's opening from the statement when the plan says so. The value
    is never re-derived by the controller: ``opening_balance`` is an ordinary editable field,
    only ``carry_forward_opening`` writes it, and that runs once, in ``create_for_week``."""
    old = flt(doc.opening_balance, 2)
    if wk.get("opening_action") != "set":
        return _opening_result(wk, old=old)
    new = flt(wk["opening_statement"], 2)
    doc.opening_balance = new
    res = _opening_result(wk, old=old, new=new)
    res["difference"] = round(new - old, 2)
    res["note"] = f"Opening set from statement: {new:,.2f} (system carry-forward was {old:,.2f}, difference {new - old:,.2f})"
    return res


def _set_opening_only(wk):
    """A week with nothing new to import whose open sheet disagrees with the statement."""
    doc = _open_sheet(wk["sheet"], getdate(wk["week_ending"]))
    res = _apply_opening(doc, wk)
    doc.save()
    return res
